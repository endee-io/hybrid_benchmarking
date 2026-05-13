#!/usr/bin/env python3
"""
Qdrant Hybrid Benchmark Script — Splade, Concurrency Sweep
Runs query-only benchmarks across concurrency values: 2, 4, 8, 16, 24
Each concurrency is run 3 times; best QPS result is recorded.
Output: Excel file with QPS, recall, latency, ndcg, map per concurrency.
"""

import json
import os
import subprocess
import time
from datetime import datetime

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

# ============================================================
# CONFIGURATION
# ============================================================
HOST             = "148.113.58.83"
INDEX_NAME       = "beir_quora_splade_int8"
DATASET_NAME     = "beir_quora"
SPARSE_MODE      = "splade"
DATA_DIR         = "data"
TOP_K            = 30
CACHE_DIR        = "model_cache"
VALIDATION_VENV  = "validation-env"
DATATYPE         = "int8"   # float32, float16, int8
PREFER_GRPC      = True
DENSE_MODEL      = "sentence-transformers/all-MiniLM-L6-v2 (384 dim)"
SPARSE_MODEL     = "Splade_PP_en_v1"

CONCURRENCY_VALUES    = [2, 4, 8, 16, 24]
RUNS_PER_CONCURRENCY  = 3
WAIT_BETWEEN_RUNS     = 20   # seconds
WAIT_BETWEEN_CONFIGS  = 20   # seconds

WORK_DIR = os.path.dirname(os.path.abspath(__file__))

OUTPUT_EXCEL = os.path.join(
    WORK_DIR,
    f"bench_qdrant_splade_concurrency_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
)


# ============================================================
# BENCHMARK RUNNER
# ============================================================

def results_label(concurrency: int, iteration: int) -> str:
    return f"{INDEX_NAME}_c{concurrency}_it{iteration}"


def results_dir(concurrency: int, iteration: int) -> str:
    label = results_label(concurrency, iteration)
    return os.path.join(WORK_DIR, "results", "qdrant", f"{label}_concurrency{concurrency}")


def run_once(concurrency: int, iteration: int) -> dict:
    label = results_label(concurrency, iteration)
    cmd = [
        "python3", "-m", "src.main",
        "--db", "qdrant",
        "--host", HOST,
        "--index-name", INDEX_NAME,
        "--dataset-name", DATASET_NAME,
        "--sparse-mode", SPARSE_MODE,
        "--data-dir", DATA_DIR,
        "--concurrency", str(concurrency),
        "--top-k", str(TOP_K),
        "--results", label,
        "--datatype", DATATYPE,
        "--cache-dir", CACHE_DIR,
        "--validation-venv", VALIDATION_VENV,
        "--skip-indexing",
    ]
    if PREFER_GRPC:
        cmd.append("--prefer-grpc")

    proc = subprocess.run(cmd, text=True, cwd=WORK_DIR)
    if proc.returncode != 0:
        print(f"  [WARN] Command exited with code {proc.returncode}")

    rdir             = results_dir(concurrency, iteration)
    summary_path     = os.path.join(rdir, "summary.json")
    correctness_path = os.path.join(rdir, "correctness.json")
    p99_path         = os.path.join(rdir, "p99_latency.json")

    if not os.path.exists(summary_path):
        print(f"  [ERROR] summary.json not found at {rdir}")
        return {"qps": None, "p99_latency_ms": None, "ndcg": None, "recall": None, "map": None}

    with open(summary_path) as f:
        summary = json.load(f)

    p99_data = {}
    if os.path.exists(p99_path):
        with open(p99_path) as f:
            p99_data = json.load(f)

    correctness = {}
    if os.path.exists(correctness_path):
        with open(correctness_path) as f:
            correctness = json.load(f)

    qps    = summary.get("qps")
    p99    = p99_data.get("p99_latency_ms")
    ndcg   = correctness.get(f"ndcg@{TOP_K}")
    recall = correctness.get(f"recall@{TOP_K}")
    map_   = correctness.get(f"map@{TOP_K}")

    print(f"  [METRICS] qps={qps}, p99={p99}ms, ndcg={ndcg}, recall={recall}, map={map_}")
    return {"qps": qps, "p99_latency_ms": p99, "ndcg": ndcg, "recall": recall, "map": map_}


def run_best_of_n(concurrency: int) -> dict:
    results = []
    for iteration in range(1, RUNS_PER_CONCURRENCY + 1):
        print(f"\n  [ITERATION {iteration}/{RUNS_PER_CONCURRENCY}] concurrency={concurrency}")
        result = run_once(concurrency, iteration)
        results.append(result)
        if iteration < RUNS_PER_CONCURRENCY:
            print(f"  [WAIT] {WAIT_BETWEEN_RUNS}s before next iteration ...")
            time.sleep(WAIT_BETWEEN_RUNS)

    best = max(results, key=lambda r: r.get("qps") or 0)
    best_iter = results.index(best) + 1
    print(f"  [BEST] Iteration {best_iter}: qps={best.get('qps')}")
    return best


# ============================================================
# EXCEL WRITER
# ============================================================

def write_excel(rows: list, output_path: str):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Qdrant Splade Concurrency"

    thin   = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left   = Alignment(horizontal="left",   vertical="center", wrap_text=True)

    HDR_BG   = "1B4F72"
    HDR_FONT = Font(bold=True, color="FFFFFF")
    ROW_ODD  = "EAF2F8"
    ROW_EVEN = "FFFFFF"

    columns = [
        ("Dataset",              24),
        ("Dense Model",          30),
        ("Sparse Model",         28),
        ("Datatype",             12),
        ("top-k",                 8),
        ("Concurrency",          14),
        ("Recall (%)",           12),
        ("QPS",                  14),
        ("Latency p99 (ms)",     18),
        ("NDCG",                 12),
        ("MAP",                  12),
    ]

    NUM_COLS = len(columns)

    for col_idx, (_, width) in enumerate(columns, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = width

    current_row = 1

    title_cell = ws.cell(
        row=current_row, column=1,
        value=f"Qdrant Splade Benchmark — {INDEX_NAME} (best of {RUNS_PER_CONCURRENCY} runs, top-k={TOP_K})"
    )
    title_cell.font      = Font(bold=True, size=12)
    title_cell.alignment = left
    ws.merge_cells(start_row=current_row, start_column=1,
                   end_row=current_row,   end_column=NUM_COLS)
    ws.row_dimensions[current_row].height = 22
    current_row += 1

    ws.row_dimensions[current_row].height = 28
    for col_idx, (header, _) in enumerate(columns, start=1):
        c = ws.cell(row=current_row, column=col_idx, value=header)
        c.font      = HDR_FONT
        c.fill      = PatternFill("solid", fgColor=HDR_BG)
        c.alignment = center
        c.border    = border
    current_row += 1

    for i, r in enumerate(rows):
        ws.row_dimensions[current_row].height = 22
        bg = ROW_ODD if i % 2 == 0 else ROW_EVEN
        rf = PatternFill("solid", fgColor=bg)

        qps    = r.get("qps")
        p99    = r.get("p99_latency_ms")
        ndcg   = r.get("ndcg")
        recall = r.get("recall")
        map_   = r.get("map")

        values = [
            DATASET_NAME,
            DENSE_MODEL,
            SPARSE_MODEL,
            DATATYPE,
            TOP_K,
            r["concurrency"],
            round(recall * 100, 3) if recall is not None else "N/A",
            round(qps,    4)       if qps    is not None else "N/A",
            round(p99,    4)       if p99    is not None else "N/A",
            round(ndcg  * 100, 3) if ndcg   is not None else "N/A",
            round(map_  * 100, 3) if map_   is not None else "N/A",
        ]

        for col_idx, val in enumerate(values, start=1):
            c = ws.cell(row=current_row, column=col_idx, value=val)
            c.fill      = rf
            c.alignment = left if col_idx <= 3 else center
            c.border    = border

        current_row += 1

    wb.save(output_path)
    print(f"\n[EXCEL] Saved → {output_path}")


# ============================================================
# MAIN
# ============================================================

def main():
    print("=" * 60)
    print("Qdrant Splade Benchmark — Concurrency Sweep")
    print(f"Index        : {INDEX_NAME}")
    print(f"Datatype     : {DATATYPE}")
    print(f"gRPC         : {PREFER_GRPC}")
    print(f"top-k        : {TOP_K}")
    print(f"Concurrency  : {CONCURRENCY_VALUES}")
    print(f"Runs/config  : {RUNS_PER_CONCURRENCY}")
    print(f"Output       : {OUTPUT_EXCEL}")
    print("=" * 60)

    rows = []

    for concurrency in CONCURRENCY_VALUES:
        print(f"\n{'='*40}")
        print(f"  concurrency = {concurrency}")
        print(f"{'='*40}")
        metrics = run_best_of_n(concurrency)
        rows.append({"concurrency": concurrency, **metrics})
        print(f"\n  [WAIT] {WAIT_BETWEEN_CONFIGS}s before next concurrency ...")
        time.sleep(WAIT_BETWEEN_CONFIGS)

    write_excel(rows, OUTPUT_EXCEL)


if __name__ == "__main__":
    main()
