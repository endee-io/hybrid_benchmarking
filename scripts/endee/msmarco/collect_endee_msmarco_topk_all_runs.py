#!/usr/bin/env python3
"""
Collect all top-k benchmark runs from the results folder into a single Excel sheet.
Reads every iteration (not just best) for a given collection name.
Best run per top-k is highlighted in green.
A thick blue separator line is drawn after each top-k group.
"""

import json
import os
import re
from datetime import datetime

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

# ============================================================
# CONFIGURATION — set to match the collection you ran
# ============================================================
COLLECTION_NAME = "beir_msmarco_int16"
DATASET_NAME    = "beir_msmarco"
DENSE_MODEL     = "sentence-transformers/all-MiniLM-L6-v2 (384 dim)"
SPARSE_MODEL    = "PyMilvus BM25EmbeddingFunction"
PRECISION       = "int16"

SCRIPT_DIR   = os.path.dirname(os.path.abspath(__file__))
PROJECT_ROOT = os.path.abspath(os.path.join(SCRIPT_DIR, "../../.."))
RESULTS_BASE = os.path.join(PROJECT_ROOT, "results", "endee")

OUTPUT_EXCEL = os.path.join(
    SCRIPT_DIR,
    f"collect_{COLLECTION_NAME}_topk_all_runs_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
)

# Matches: {COLLECTION_NAME}_t{top_k}_it{iteration}_concurrency{N}
FOLDER_RE = re.compile(r"^(.+)_t(\d+)_it(\d+)_concurrency(\d+)$")


# ============================================================
# COLLECT
# ============================================================

def load_run(folder_path: str, top_k: int, iteration: int, concurrency: int) -> dict | None:
    summary_path     = os.path.join(folder_path, "summary.json")
    correctness_path = os.path.join(folder_path, "correctness.json")
    p99_path         = os.path.join(folder_path, "p99_latency.json")

    if not os.path.exists(summary_path):
        return None

    with open(summary_path) as f:
        summary = json.load(f)

    correctness = {}
    if os.path.exists(correctness_path):
        with open(correctness_path) as f:
            correctness = json.load(f)

    p99_data = {}
    if os.path.exists(p99_path):
        with open(p99_path) as f:
            p99_data = json.load(f)

    return {
        "top_k":       top_k,
        "iteration":   iteration,
        "concurrency": concurrency,
        "qps":         summary.get("qps"),
        "p99":         p99_data.get("p99_latency_ms") or summary.get("p99_latency_ms"),
        "ndcg":        correctness.get(f"ndcg@{top_k}"),
        "recall":      correctness.get(f"recall@{top_k}"),
        "map":         correctness.get(f"map@{top_k}"),
    }


def collect_rows() -> list:
    rows = []
    if not os.path.isdir(RESULTS_BASE):
        print(f"[ERROR] Results directory not found: {RESULTS_BASE}")
        return rows

    for folder_name in os.listdir(RESULTS_BASE):
        m = FOLDER_RE.match(folder_name)
        if not m:
            continue
        if m.group(1) != COLLECTION_NAME:
            continue

        top_k       = int(m.group(2))
        iteration   = int(m.group(3))
        concurrency = int(m.group(4))
        folder_path = os.path.join(RESULTS_BASE, folder_name)

        run = load_run(folder_path, top_k, iteration, concurrency)
        if run:
            rows.append(run)

    rows.sort(key=lambda r: (r["top_k"], r["iteration"]))
    return rows


def mark_best(rows: list) -> set:
    """Return (top_k, iteration) pairs with best QPS per top_k group."""
    by_topk = {}
    for r in rows:
        t = r["top_k"]
        if t not in by_topk:
            by_topk[t] = r
        elif (r["qps"] or 0) > (by_topk[t]["qps"] or 0):
            by_topk[t] = r
    return {(r["top_k"], r["iteration"]) for r in by_topk.values()}


# ============================================================
# EXCEL WRITER
# ============================================================

def write_excel(rows: list, output_path: str):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "All Runs"

    thin       = Side(style="thin",   color="000000")
    thick_blue = Side(style="medium", color="1B4F72")
    border     = Border(left=thin, right=thin, top=thin, bottom=thin)
    center     = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left       = Alignment(horizontal="left",   vertical="center", wrap_text=True)

    HDR_BG   = "1B4F72"
    HDR_FONT = Font(bold=True, color="FFFFFF")
    BEST_BG  = "D5F5E3"   # light green — best QPS run per top-k
    ROW_ODD  = "EAF2F8"
    ROW_EVEN = "FFFFFF"

    columns = [
        ("Dataset",           24),
        ("Dense Model",       30),
        ("Sparse Model",      28),
        ("Precision",         12),
        ("top-k",              8),
        ("Concurrency",       14),
        ("Iteration",         10),
        ("QPS",               14),
        ("Latency p99 (ms)",  18),
        ("Recall (%)",        12),
        ("NDCG",              12),
        ("MAP",               12),
        ("Best Run?",         12),
    ]

    NUM_COLS = len(columns)

    for col_idx, (_, width) in enumerate(columns, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = width

    current_row = 1

    title_cell = ws.cell(
        row=current_row, column=1,
        value=f"Endee Top-K Benchmark — {COLLECTION_NAME} — All Runs"
    )
    title_cell.font      = Font(bold=True, size=12)
    title_cell.alignment = left
    ws.merge_cells(start_row=current_row, start_column=1,
                   end_row=current_row,   end_column=NUM_COLS)
    ws.row_dimensions[current_row].height = 22
    current_row += 1

    ws.row_dimensions[current_row].height = 28
    for col_idx, (header, _) in enumerate(columns, start=1):
        c = ws.cell(row=current_row, column=col_idx, value=header)
        c.font      = HDR_FONT
        c.fill      = PatternFill("solid", fgColor=HDR_BG)
        c.alignment = center
        c.border    = border
    current_row += 1

    best = mark_best(rows)

    for i, r in enumerate(rows):
        is_best    = (r["top_k"], r["iteration"]) in best
        is_last_in_group = (
            i == len(rows) - 1 or rows[i + 1]["top_k"] != r["top_k"]
        )

        ws.row_dimensions[current_row].height = 22
        bg = BEST_BG if is_best else (ROW_ODD if i % 2 == 0 else ROW_EVEN)
        rf = PatternFill("solid", fgColor=bg)

        bottom_side = thick_blue if is_last_in_group else thin

        qps    = r.get("qps")
        p99    = r.get("p99")
        ndcg   = r.get("ndcg")
        recall = r.get("recall")
        map_   = r.get("map")

        values = [
            DATASET_NAME,
            DENSE_MODEL,
            SPARSE_MODEL,
            PRECISION,
            r["top_k"],
            r["concurrency"],
            r["iteration"],
            round(qps,    4)       if qps    is not None else "N/A",
            round(p99,    4)       if p99    is not None else "N/A",
            round(recall * 100, 3) if recall is not None else "N/A",
            round(ndcg   * 100, 3) if ndcg   is not None else "N/A",
            round(map_   * 100, 3) if map_   is not None else "N/A",
            "YES" if is_best else "",
        ]

        for col_idx, val in enumerate(values, start=1):
            c = ws.cell(row=current_row, column=col_idx, value=val)
            c.fill      = rf
            c.alignment = left if col_idx <= 3 else center
            c.border    = Border(
                left=thin, right=thin, top=thin, bottom=bottom_side
            )

        current_row += 1

    wb.save(output_path)
    print(f"\n[EXCEL] Saved → {output_path}")


# ============================================================
# MAIN
# ============================================================

def main():
    print("=" * 60)
    print(f"Collecting all top-k runs for: {COLLECTION_NAME}")
    print(f"Results base: {RESULTS_BASE}")
    print("=" * 60)

    rows = collect_rows()

    if not rows:
        print("[WARN] No matching result folders found.")
        return

    top_k_groups = sorted({r["top_k"] for r in rows})
    print(f"Found {len(rows)} run(s) across top-k values: {top_k_groups}")
    write_excel(rows, OUTPUT_EXCEL)


if __name__ == "__main__":
    main()
