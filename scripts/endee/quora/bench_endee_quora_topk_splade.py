#!/usr/bin/env python3
"""
Endee Hybrid Benchmark Script — Splade, Top-K Sweep
Runs query-only benchmarks across top-k values.
Each top-k is run 3 times; best QPS result is recorded.
Output: Excel file with QPS, recall, latency, ndcg, map per top-k.
"""

import json
import os
import subprocess
import time
from datetime import datetime

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

# ============================================================
# CONFIGURATION
# ============================================================
BASE_URL         = "http://148.113.58.83:8080/api/v2"
VECTOR_TOKEN     = "mytoken"
COLLECTION_NAME  = "beir_quora_splade_float16"
DATASET_NAME     = "beir_quora"
SPARSE_MODE      = "splade"
DATA_DIR         = "data"
CONCURRENCY      = 16
CACHE_DIR        = "model_cache"
VALIDATION_VENV  = "validation-env"
PRECISION        = "float16"
DENSE_MODEL      = "sentence-transformers/all-MiniLM-L6-v2 (384 dim)"
SPARSE_MODEL     = "Splade_PP_en_v1"

TOP_K_VALUES      = [10, 30, 60, 100, 500, 1000]
RUNS_PER_TOP_K    = 3
WAIT_BETWEEN_RUNS = 20   # seconds
WAIT_BETWEEN_TOPK = 20   # seconds

WORK_DIR = os.path.dirname(os.path.abspath(__file__))

OUTPUT_EXCEL = os.path.join(
    WORK_DIR,
    f"bench_endee_splade_topk_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
)


# ============================================================
# BENCHMARK RUNNER
# ============================================================

def results_label(top_k: int, iteration: int) -> str:
    return f"{COLLECTION_NAME}_t{top_k}_it{iteration}"


def results_dir(top_k: int, iteration: int) -> str:
    label = results_label(top_k, iteration)
    return os.path.join(WORK_DIR, "results", "endee", f"{label}_concurrency{CONCURRENCY}")


def run_once(top_k: int, iteration: int) -> dict:
    label = results_label(top_k, iteration)
    cmd = [
        "python3", "-m", "src.main",
        "--db", "endee",
        "--base-url", BASE_URL,
        "--vector-token", VECTOR_TOKEN,
        "--index-name", COLLECTION_NAME,
        "--dataset-name", DATASET_NAME,
        "--sparse-mode", SPARSE_MODE,
        "--data-dir", DATA_DIR,
        "--concurrency", str(CONCURRENCY),
        "--top-k", str(top_k),
        "--results", label,
        "--precision", PRECISION,
        "--cache-dir", CACHE_DIR,
        "--validation-venv", VALIDATION_VENV,
        "--skip-indexing",
    ]

    proc = subprocess.run(cmd, text=True, cwd=WORK_DIR)
    if proc.returncode != 0:
        print(f"  [WARN] Command exited with code {proc.returncode}")

    rdir             = results_dir(top_k, iteration)
    summary_path     = os.path.join(rdir, "summary.json")
    correctness_path = os.path.join(rdir, "correctness.json")
    p99_path         = os.path.join(rdir, "p99_latency.json")

    if not os.path.exists(summary_path):
        print(f"  [ERROR] summary.json not found at {rdir}")
        return {"qps": None, "p99_latency_ms": None, "ndcg": None, "recall": None, "map": None}

    with open(summary_path) as f:
        summary = json.load(f)

    p99_data = {}
    if os.path.exists(p99_path):
        with open(p99_path) as f:
            p99_data = json.load(f)

    correctness = {}
    if os.path.exists(correctness_path):
        with open(correctness_path) as f:
            correctness = json.load(f)

    qps    = summary.get("qps")
    p99    = p99_data.get("p99_latency_ms")
    ndcg   = correctness.get(f"ndcg@{top_k}")
    recall = correctness.get(f"recall@{top_k}")
    map_   = correctness.get(f"map@{top_k}")

    print(f"  [METRICS] qps={qps}, p99={p99}ms, ndcg={ndcg}, recall={recall}, map={map_}")
    return {"qps": qps, "p99_latency_ms": p99, "ndcg": ndcg, "recall": recall, "map": map_}


def run_best_of_n(top_k: int) -> dict:
    results = []
    for iteration in range(1, RUNS_PER_TOP_K + 1):
        print(f"\n  [ITERATION {iteration}/{RUNS_PER_TOP_K}] top_k={top_k}")
        result = run_once(top_k, iteration)
        results.append(result)
        if iteration < RUNS_PER_TOP_K:
            print(f"  [WAIT] {WAIT_BETWEEN_RUNS}s before next iteration ...")
            time.sleep(WAIT_BETWEEN_RUNS)

    best = max(results, key=lambda r: r.get("qps") or 0)
    best_iter = results.index(best) + 1
    print(f"  [BEST] Iteration {best_iter}: qps={best.get('qps')}")
    return best


# ============================================================
# EXCEL WRITER
# ============================================================

def write_excel(rows: list, output_path: str):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Endee Splade Top-K"

    thin   = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left   = Alignment(horizontal="left",   vertical="center", wrap_text=True)

    HDR_BG   = "1B4F72"
    HDR_FONT = Font(bold=True, color="FFFFFF")
    ROW_ODD  = "EAF2F8"
    ROW_EVEN = "FFFFFF"

    columns = [
        ("Dataset",              24),
        ("Dense Model",          30),
        ("Sparse Model",         28),
        ("Precision",            12),
        ("Concurrency",          14),
        ("top-k",                 8),
        ("Recall (%)",           12),
        ("QPS",                  14),
        ("Latency p99 (ms)",     18),
        ("NDCG",                 12),
        ("MAP",                  12),
    ]

    NUM_COLS = len(columns)

    for col_idx, (_, width) in enumerate(columns, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = width

    current_row = 1

    title_cell = ws.cell(
        row=current_row, column=1,
        value=f"Endee Splade Benchmark — {COLLECTION_NAME} (best of {RUNS_PER_TOP_K} runs)"
    )
    title_cell.font      = Font(bold=True, size=12)
    title_cell.alignment = left
    ws.merge_cells(start_row=current_row, start_column=1,
                   end_row=current_row,   end_column=NUM_COLS)
    ws.row_dimensions[current_row].height = 22
    current_row += 1

    ws.row_dimensions[current_row].height = 28
    for col_idx, (header, _) in enumerate(columns, start=1):
        c = ws.cell(row=current_row, column=col_idx, value=header)
        c.font      = HDR_FONT
        c.fill      = PatternFill("solid", fgColor=HDR_BG)
        c.alignment = center
        c.border    = border
    current_row += 1

    for i, r in enumerate(rows):
        ws.row_dimensions[current_row].height = 22
        bg = ROW_ODD if i % 2 == 0 else ROW_EVEN
        rf = PatternFill("solid", fgColor=bg)

        qps    = r.get("qps")
        p99    = r.get("p99_latency_ms")
        ndcg   = r.get("ndcg")
        recall = r.get("recall")
        map_   = r.get("map")

        values = [
            DATASET_NAME,
            DENSE_MODEL,
            SPARSE_MODEL,
            PRECISION,
            CONCURRENCY,
            r["top_k"],
            round(recall * 100, 3) if recall is not None else "N/A",
            round(qps,    4)       if qps    is not None else "N/A",
            round(p99,    4)       if p99    is not None else "N/A",
            round(ndcg  * 100, 3) if ndcg   is not None else "N/A",
            round(map_  * 100, 3) if map_   is not None else "N/A",
        ]

        for col_idx, val in enumerate(values, start=1):
            c = ws.cell(row=current_row, column=col_idx, value=val)
            c.fill      = rf
            c.alignment = left if col_idx <= 3 else center
            c.border    = border

        current_row += 1

    wb.save(output_path)
    print(f"\n[EXCEL] Saved → {output_path}")


# ============================================================
# MAIN
# ============================================================

def main():
    print("=" * 60)
    print("Endee Splade Benchmark — Top-K Sweep")
    print(f"Collection: {COLLECTION_NAME}")
    print(f"Precision  : {PRECISION}")
    print(f"Top-K      : {TOP_K_VALUES}")
    print(f"Runs/top-k : {RUNS_PER_TOP_K}")
    print(f"Output     : {OUTPUT_EXCEL}")
    print("=" * 60)

    rows = []

    for top_k in TOP_K_VALUES:
        print(f"\n{'='*40}")
        print(f"  top_k = {top_k}")
        print(f"{'='*40}")
        metrics = run_best_of_n(top_k)
        rows.append({"top_k": top_k, **metrics})
        print(f"\n  [WAIT] {WAIT_BETWEEN_TOPK}s before next top-k ...")
        time.sleep(WAIT_BETWEEN_TOPK)

    write_excel(rows, OUTPUT_EXCEL)


if __name__ == "__main__":
    main()
