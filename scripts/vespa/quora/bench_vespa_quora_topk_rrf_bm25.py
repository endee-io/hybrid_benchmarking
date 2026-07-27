#!/usr/bin/env python3
"""
Vespa rrf_bm25_text Benchmark Script
Runs query-only benchmarks across top-k values: 10, 30, 60, 100, 500, 1000
Each top-k is run 4 times; best QPS result is recorded.
Output: Excel file with QPS, recall, latency per top-k.
"""

import json
import os
import subprocess
import time
from datetime import datetime

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

# ============================================================
# CONFIGURATION
# ============================================================
VESPA_URL        = "http://148.113.58.83"
VESPA_PORT       = 8080
VESPA_CONFIG_PORT = 19071
QUERY_MODE       = "rrf_bm25_text"
EF_CONSTRUCTION  = 128
EF_SEARCH        = 128
PRECISION        = "int8"
INDEX_NAME       = "vespa_hybrid_int8_rerank50"
DATASET_NAME     = "beir_quora"
SPARSE_MODE      = "bm25"
HF_DATASET_ID    = "BeIR/quora"
DATA_DIR         = "data"
CONCURRENCY      = 16
RERANK_COUNT     = 50
CACHE_DIR        = "model_cache"
VALIDATION_VENV  = "validation-env"

TOP_K_VALUES     = [10, 30, 60, 100, 500, 1000]
RESULTS_BASE     = "results/vespa"
RUNS_PER_TOP_K   = 4
WAIT_BETWEEN_RUNS = 20   # seconds
WAIT_BETWEEN_TOPK = 20  # seconds

OUTPUT_EXCEL = os.path.join(
    "/home/debian/hybrid_benchmarking",
    f"bench_vespa_topk_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
)

# ============================================================
# BENCHMARK RUNNER
# ============================================================

def results_dir(top_k: int) -> str:
    label = f"{INDEX_NAME}_rrf_bm25text_{PRECISION}_topk{top_k}"
    return os.path.join(RESULTS_BASE, f"{label}_concurrency{CONCURRENCY}")


def run_once(top_k: int) -> dict:
    label = f"{INDEX_NAME}_rrf_bm25text_{PRECISION}_topk{top_k}"
    cmd = [
        "python3.13", "-m", "src.main",
        "--db", "vespa",
        "--vespa-url", VESPA_URL,
        "--vespa-port", str(VESPA_PORT),
        "--vespa-config-port", str(VESPA_CONFIG_PORT),
        "--vespa-query-mode", QUERY_MODE,
        "--vespa-ef-construction", str(EF_CONSTRUCTION),
        "--vespa-ef-search", str(EF_SEARCH),
        "--vespa-precision", PRECISION,
        "--vespa-rerank-count", str(RERANK_COUNT),
        "--index-name", INDEX_NAME,
        "--dataset-name", DATASET_NAME,
        "--sparse-mode", SPARSE_MODE,
        "--hf-dataset-id", HF_DATASET_ID,
        "--data-dir", DATA_DIR,
        "--concurrency", str(CONCURRENCY),
        "--top-k", str(top_k),
        "--results", label,
        "--cache-dir", CACHE_DIR,
        "--skip-indexing",
        "--validation-venv", VALIDATION_VENV,
    ]

    proc = subprocess.run(cmd, text=True, cwd="/home/debian/hybrid_benchmarking")
    if proc.returncode != 0:
        print(f"  [WARN] Command exited with code {proc.returncode}")

    rdir = results_dir(top_k)
    summary_path     = os.path.join(rdir, "summary.json")
    correctness_path = os.path.join(rdir, "correctness.json")

    if not os.path.exists(summary_path):
        print(f"  [ERROR] summary.json not found at {rdir}")
        return {"qps": None, "p99_latency_ms": None, "ndcg": None, "recall": None}

    with open(summary_path) as f:
        summary = json.load(f)
    with open(correctness_path) as f:
        correctness = json.load(f)

    qps     = summary.get("qps")
    p99     = summary.get("p99_latency_ms")
    ndcg    = correctness.get(f"ndcg@{top_k}")
    recall  = correctness.get(f"recall@{top_k}")

    print(f"  [METRICS] qps={qps}, p99={p99}ms, ndcg={ndcg}, recall={recall}")
    return {"qps": qps, "p99_latency_ms": p99, "ndcg": ndcg, "recall": recall}


def run_best_of_n(top_k: int) -> dict:
    results = []
    for attempt in range(1, RUNS_PER_TOP_K + 1):
        print(f"\n  [ATTEMPT {attempt}/{RUNS_PER_TOP_K}] top_k={top_k}")
        result = run_once(top_k)
        results.append(result)
        if attempt < RUNS_PER_TOP_K:
            print(f"  [WAIT] {WAIT_BETWEEN_RUNS}s before next attempt ...")
            time.sleep(WAIT_BETWEEN_RUNS)

    best = max(results, key=lambda r: r.get("qps") or 0)
    best_attempt = results.index(best) + 1
    print(f"  [BEST] Attempt {best_attempt}: qps={best.get('qps')}")
    return best


# ============================================================
# EXCEL WRITER
# ============================================================

def write_excel(rows: list, output_path: str):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Vespa rrf_bm25_text"

    thin   = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left   = Alignment(horizontal="left", vertical="center")

    HDR_BG   = "1B4F72"
    HDR_FONT = Font(bold=True, color="FFFFFF")
    ROW_ODD  = "EAF2F8"
    ROW_EVEN = "FFFFFF"

    columns = [
        ("Dataset",              20),
        ("Index",                18),
        ("Precision",            12),
        ("Query Mode",           18),
        ("ef_construction",      16),
        ("Rerank Count",         14),
        ("Concurrency",          14),
        ("top-k",                 8),
        ("QPS",                  12),
        ("Latency p99 (ms)",     18),
        ("NDCG",                 12),
        ("Recall (%)",           12),
    ]

    NUM_COLS = len(columns)

    for col_idx, (_, width) in enumerate(columns, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = width

    current_row = 1

    # Title
    title_cell = ws.cell(
        row=current_row, column=1,
        value=f"Vespa {QUERY_MODE} Benchmark — {INDEX_NAME} (best of {RUNS_PER_TOP_K} runs)"
    )
    title_cell.font      = Font(bold=True, size=12)
    title_cell.alignment = left
    ws.merge_cells(start_row=current_row, start_column=1,
                   end_row=current_row,   end_column=NUM_COLS)
    ws.row_dimensions[current_row].height = 22
    current_row += 1

    # Header
    ws.row_dimensions[current_row].height = 28
    for col_idx, (header, _) in enumerate(columns, start=1):
        c = ws.cell(row=current_row, column=col_idx, value=header)
        c.font      = HDR_FONT
        c.fill      = PatternFill("solid", fgColor=HDR_BG)
        c.alignment = center
        c.border    = border
    current_row += 1

    # Data rows
    for i, r in enumerate(rows):
        ws.row_dimensions[current_row].height = 22
        bg = ROW_ODD if i % 2 == 0 else ROW_EVEN
        rf = PatternFill("solid", fgColor=bg)

        qps    = r.get("qps")
        p99    = r.get("p99_latency_ms")
        ndcg   = r.get("ndcg")
        recall = r.get("recall")

        values = [
            DATASET_NAME,
            INDEX_NAME,
            PRECISION,
            QUERY_MODE,
            EF_CONSTRUCTION,
            RERANK_COUNT,
            CONCURRENCY,
            r["top_k"],
            round(qps,    4)       if qps    is not None else "N/A",
            round(p99,    4)       if p99    is not None else "N/A",
            round(ndcg,   4)       if ndcg   is not None else "N/A",
            round(recall * 100, 2) if recall is not None else "N/A",
        ]

        for col_idx, val in enumerate(values, start=1):
            c = ws.cell(row=current_row, column=col_idx, value=val)
            c.fill      = rf
            c.alignment = left if col_idx == 1 else center
            c.border    = border

        current_row += 1

    wb.save(output_path)
    print(f"\n[EXCEL] Saved → {output_path}")


# ============================================================
# MAIN
# ============================================================

def main():
    print("=" * 60)
    print("Vespa rrf_bm25_text Benchmark")
    print(f"Index      : {INDEX_NAME}")
    print(f"Precision  : {PRECISION}")
    print(f"Top-K      : {TOP_K_VALUES}")
    print(f"Runs/top-k : {RUNS_PER_TOP_K}")
    print(f"Output     : {OUTPUT_EXCEL}")
    print("=" * 60)

    rows = []

    for top_k in TOP_K_VALUES:
        print(f"\n{'='*40}")
        print(f"  top_k = {top_k}")
        print(f"{'='*40}")
        metrics = run_best_of_n(top_k)
        rows.append({"top_k": top_k, **metrics})
        print(f"\n  [WAIT] {WAIT_BETWEEN_TOPK}s before next top-k ...")
        time.sleep(WAIT_BETWEEN_TOPK)

    write_excel(rows, OUTPUT_EXCEL)


if __name__ == "__main__":
    main()
